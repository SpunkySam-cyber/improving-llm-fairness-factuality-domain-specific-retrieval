from __future__ import annotations

import argparse
import json
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


VERSION_DIR = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = (
    VERSION_DIR
    / "outputs"
    / "019fe4c8-40ac-78f1-88da-2d2c98dd9d55"
    / "SmartLabs_Corpus_Pilot_50_Terms.xlsx"
)
DEFAULT_REPORT = VERSION_DIR / "reports" / "corpus_pilot_validation.json"
TARGET = 50

HEADERS = {
    "Terms": [
        "Term ID", "Canonical Text", "Category", "Language Code",
        "Definition / Meaning", "Notes", "Status",
    ],
    "Variants": [
        "Variant ID", "Term ID", "Variant Text", "Variant Type",
        "Verification Status", "Notes",
    ],
    "Evidence": [
        "Evidence ID", "Term ID", "Variant ID (optional)", "Video ID",
        "Video Title", "Video URL", "Start Time (seconds)",
        "End Time (seconds)", "Context Excerpt", "Evidence Kind",
        "Review Status", "Reviewed By", "Reviewed Date", "Notes",
    ],
}
TERM_CATEGORIES = {
    "islamic_term", "quranic_phrase", "person_name", "book_title",
    "group_name", "honorific", "place_name", "other",
}
TERM_STATUSES = {"draft", "verified", "retired"}
VARIANT_TYPES = {
    "accepted_spelling", "observed_asr_error", "generated_possible_error",
}
REVIEW_STATUSES = {"pending", "verified", "rejected"}
EVIDENCE_KINDS = {"spoken_audio", "on_screen_text", "published_source"}


class WorkbookValidationError(RuntimeError):
    """Raised when the workbook structure cannot be safely validated."""


def text(value: Any) -> str | None:
    if value is None:
        return None
    cleaned = str(value).strip()
    return cleaned or None


def normalized(value: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", value).casefold().split())


def valid_date(value: Any) -> bool:
    if isinstance(value, (date, datetime)):
        return True
    if value is None:
        return False
    try:
        date.fromisoformat(str(value).strip())
    except ValueError:
        return False
    return True


def seconds(value: Any, label: str, errors: list[str]) -> float | None:
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        errors.append(f"{label}: enter numeric seconds")
        return None
    if number < 0:
        errors.append(f"{label}: time cannot be negative")
        return None
    return number


def read_sheets(path: Path) -> dict[str, list[dict[str, Any]]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except (OSError, ValueError) as exc:
        raise WorkbookValidationError(f"Cannot open workbook: {path}") from exc
    data: dict[str, list[dict[str, Any]]] = {}
    try:
        for sheet_name, expected_headers in HEADERS.items():
            if sheet_name not in workbook.sheetnames:
                raise WorkbookValidationError(f"Missing required sheet: {sheet_name}")
            sheet = workbook[sheet_name]
            actual_headers = [cell.value for cell in sheet[1]][: len(expected_headers)]
            if actual_headers != expected_headers:
                raise WorkbookValidationError(f"Headers changed on sheet: {sheet_name}")
            records = []
            for row_number, values in enumerate(
                sheet.iter_rows(min_row=2, max_col=len(expected_headers), values_only=True),
                start=2,
            ):
                record = dict(zip(expected_headers, values))
                record["row"] = row_number
                records.append(record)
            data[sheet_name] = records
    finally:
        workbook.close()
    return data


def validate(path: Path) -> dict[str, Any]:
    sheets = read_sheets(path)
    errors: list[str] = []
    warnings: list[str] = []
    terms: dict[str, dict[str, Any]] = {}
    canonical_forms: set[str] = set()

    for row in sheets["Terms"]:
        canonical = text(row["Canonical Text"])
        if canonical is None:
            continue
        label = f"Terms row {row['row']}"
        term_id = text(row["Term ID"])
        category = text(row["Category"])
        status = text(row["Status"]) or "draft"
        if term_id is None or re.fullmatch(r"T\d{3,}", term_id) is None:
            errors.append(f"{label}: invalid Term ID")
            continue
        if term_id in terms:
            errors.append(f"{label}: duplicate Term ID {term_id}")
        canonical_key = normalized(canonical)
        if canonical_key in canonical_forms:
            errors.append(f"{label}: duplicate canonical term {canonical}")
        canonical_forms.add(canonical_key)
        if category not in TERM_CATEGORIES:
            errors.append(f"{label}: select a Category")
        if status not in TERM_STATUSES:
            errors.append(f"{label}: invalid Status")
        terms[term_id] = {
            "canonical": canonical,
            "status": status,
            "row": row["row"],
        }

    variants: dict[str, dict[str, Any]] = {}
    variant_keys: set[tuple[str | None, str, str | None]] = set()
    for row in sheets["Variants"]:
        variant_text = text(row["Variant Text"])
        if variant_text is None:
            continue
        label = f"Variants row {row['row']}"
        variant_id = text(row["Variant ID"])
        term_id = text(row["Term ID"])
        variant_type = text(row["Variant Type"])
        status = text(row["Verification Status"]) or "pending"
        if variant_id is None or re.fullmatch(r"V\d{3,}", variant_id) is None:
            errors.append(f"{label}: invalid Variant ID")
            continue
        if variant_id in variants:
            errors.append(f"{label}: duplicate Variant ID {variant_id}")
        if term_id not in terms:
            errors.append(f"{label}: select an entered Term ID")
        if variant_type not in VARIANT_TYPES:
            errors.append(f"{label}: select a Variant Type")
        if status not in REVIEW_STATUSES:
            errors.append(f"{label}: invalid Verification Status")
        if variant_type == "generated_possible_error" and status == "verified":
            errors.append(f"{label}: generated possibilities cannot be verified")
        key = (term_id, normalized(variant_text), variant_type)
        if key in variant_keys:
            errors.append(f"{label}: duplicate variant for the same term and type")
        variant_keys.add(key)
        variants[variant_id] = {
            "term_id": term_id,
            "type": variant_type,
            "status": status,
            "row": row["row"],
        }

    evidence_count = 0
    verified_evidence_count = 0
    verified_term_evidence: set[str] = set()
    verified_variant_evidence: set[str] = set()
    evidence_ids: set[str] = set()
    for row in sheets["Evidence"]:
        term_id = text(row["Term ID"])
        if term_id is None:
            continue
        evidence_count += 1
        label = f"Evidence row {row['row']}"
        evidence_id = text(row["Evidence ID"])
        variant_id = text(row["Variant ID (optional)"])
        video_id = text(row["Video ID"])
        video_title = text(row["Video Title"])
        video_url = text(row["Video URL"])
        context = text(row["Context Excerpt"])
        kind = text(row["Evidence Kind"])
        status = text(row["Review Status"]) or "pending"
        reviewed_by = text(row["Reviewed By"])
        start = seconds(row["Start Time (seconds)"], f"{label} start time", errors)
        end = seconds(row["End Time (seconds)"], f"{label} end time", errors)
        if evidence_id is None or re.fullmatch(r"E\d{3,}", evidence_id) is None:
            errors.append(f"{label}: invalid Evidence ID")
        elif evidence_id in evidence_ids:
            errors.append(f"{label}: duplicate Evidence ID {evidence_id}")
        evidence_ids.add(evidence_id or "")
        if term_id not in terms:
            errors.append(f"{label}: select an entered Term ID")
        if variant_id is not None:
            if variant_id not in variants:
                errors.append(f"{label}: select an entered Variant ID")
            elif variants[variant_id]["term_id"] != term_id:
                errors.append(f"{label}: Variant ID belongs to a different term")
        if not video_id or not video_title or not video_url:
            errors.append(f"{label}: Video ID, Title, and URL are required")
        elif re.match(r"^https?://", video_url) is None:
            errors.append(f"{label}: Video URL must begin with http:// or https://")
        if start is None:
            errors.append(f"{label}: Start Time is required")
        if end is not None and start is not None and end < start:
            errors.append(f"{label}: End Time cannot be before Start Time")
        if context is None:
            errors.append(f"{label}: Context Excerpt is required")
        if kind not in EVIDENCE_KINDS:
            errors.append(f"{label}: select an Evidence Kind")
        if status not in REVIEW_STATUSES:
            errors.append(f"{label}: invalid Review Status")
        if status == "verified":
            verified_evidence_count += 1
            if not reviewed_by or not valid_date(row["Reviewed Date"]):
                errors.append(f"{label}: verified evidence requires reviewer and date")
            verified_term_evidence.add(term_id)
            if variant_id:
                verified_variant_evidence.add(variant_id)

    for term_id, term in terms.items():
        if term["status"] == "verified" and term_id not in verified_term_evidence:
            errors.append(
                f"Terms row {term['row']}: verified term has no verified Evidence row"
            )
    for variant_id, variant in variants.items():
        if (
            variant["type"] == "observed_asr_error"
            and variant["status"] == "verified"
            and variant_id not in verified_variant_evidence
        ):
            errors.append(
                f"Variants row {variant['row']}: verified observed error needs variant-specific evidence"
            )

    verified_terms = sum(term["status"] == "verified" for term in terms.values())
    if verified_terms < TARGET:
        warnings.append(f"Pilot incomplete: {TARGET - verified_terms} verified terms remaining")
    return {
        "workbook": str(path.resolve()),
        "target_verified_terms": TARGET,
        "terms_entered": len(terms),
        "verified_terms": verified_terms,
        "verified_terms_remaining": max(0, TARGET - verified_terms),
        "variants_entered": len(variants),
        "evidence_rows_entered": evidence_count,
        "verified_evidence_rows": verified_evidence_count,
        "ready_for_database_import": not errors and verified_terms >= TARGET,
        "errors": errors,
        "warnings": warnings,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Validate the 50-term corpus workbook.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    report = validate(args.workbook)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if report["errors"] else 0


if __name__ == "__main__":
    raise SystemExit(main())

